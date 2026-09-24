# -*- coding: utf-8 -*-
"""用户术语映射 Excel 的预检查与确认导入。"""

from __future__ import annotations

import re
import unicodedata
import uuid
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from tortoise.transactions import in_transaction

from applications.jiayueyang.models.term_mapping_model import RagTermMappingSource
from applications.jiayueyang.services.term_mapping_crud import source_term_mapping_crud
from core.exceptions import ParameterException


TEMPLATE_COLUMNS: Tuple[str, ...] = (
    "中文标准词",
    "中文别名",
    "官方英文词",
    "英文别名、缩写",
    "适用范围",
    "参考出处或补充说明",
)
_REQUIRED_COLUMNS = {"中文标准词", "官方英文词"}
_SPLIT_RE = re.compile(r"[；;]")
_SPACE_RE = re.compile(r"\s+")
_SEPARATOR_RE = re.compile(r"[-_/\\|,，、:：]+")


def _text(value: Any) -> str:
    """把 Excel 单元格转换成稳定的可保存文本。"""
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ")
    return str(value).strip()


def normalize_term(value: Any) -> str:
    """按模板约定生成用于匹配和去重的规范化键。"""
    text = unicodedata.normalize("NFKC", _text(value)).lower()
    text = _SEPARATOR_RE.sub(" ", text)
    text = _SPACE_RE.sub(" ", text)
    return text.strip()


def split_values(value: Any) -> List[str]:
    """按中文分号拆分多值，并清理空值和重复项。"""
    text = _text(value)
    if not text:
        return []
    result: List[str] = []
    seen = set()
    for item in _SPLIT_RE.split(text):
        item = item.strip()
        if not item:
            continue
        key = normalize_term(item)
        if key and key not in seen:
            result.append(item)
            seen.add(key)
    return result


def _raw_value(value: Any) -> Any:
    """确保 raw_payload 可被 JSONField 序列化。"""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return _text(value)


class TermMappingImportService:
    """解析、校验和导入用户权威映射；不访问 learned 表。"""

    @staticmethod
    def _find_header_row(rows: List[Tuple[Any, ...]]) -> Tuple[int, Dict[str, int]]:
        for row_index, row in enumerate(rows[:20], start=1):
            headers = {_text(value): index for index, value in enumerate(row)}
            if _REQUIRED_COLUMNS.issubset(headers):
                missing = [column for column in TEMPLATE_COLUMNS if column not in headers]
                if missing:
                    raise ParameterException(message=f"模板缺少固定列: {', '.join(missing)}")
                return row_index, {column: headers[column] for column in TEMPLATE_COLUMNS}
        raise ParameterException(message=f"未找到模板表头，必须包含: {'、'.join(TEMPLATE_COLUMNS)}")

    @staticmethod
    def _check_suffix(filename: str) -> None:
        if not filename.lower().endswith((".xlsx", ".xlsm")):
            raise ParameterException(message="只支持 .xlsx 或 .xlsm 格式的映射模板")

    def parse(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        self._check_suffix(filename)
        if not file_bytes:
            raise ParameterException(message="上传文件为空")
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise ParameterException(message=f"Excel 文件无法读取: {exc}") from exc

        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            header_row, column_indexes = self._find_header_row(rows)
            parsed: List[Dict[str, Any]] = []
            skipped = 0
            errors: List[Dict[str, Any]] = []
            conflicts: List[Dict[str, Any]] = []
            key_rows: Dict[str, int] = {}
            alias_targets: Dict[Tuple[str, str], Tuple[str, int]] = {}

            for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
                raw_payload = {
                    column: _raw_value(row[index] if index < len(row) else None)
                    for column, index in column_indexes.items()
                }
                if not any(_text(value) for value in raw_payload.values()):
                    skipped += 1
                    continue

                canonical = _text(raw_payload["中文标准词"])
                scope = _text(raw_payload["适用范围"])
                aliases_zh = split_values(raw_payload["中文别名"])
                aliases_en = split_values(raw_payload["英文别名、缩写"])
                official_en = _text(raw_payload["官方英文词"]) or None
                row_errors: List[str] = []
                if not canonical:
                    row_errors.append("中文标准词不能为空")

                original_alias_values = [
                    item.strip() for item in _SPLIT_RE.split(_text(raw_payload["中文别名"])) if item.strip()
                ] + [
                    item.strip() for item in _SPLIT_RE.split(_text(raw_payload["英文别名、缩写"])) if item.strip()
                ]
                normalized_alias_values = [normalize_term(item) for item in original_alias_values]
                if len(normalized_alias_values) != len(set(normalized_alias_values)):
                    row_errors.append("别名重复，请合并后只保留一个")

                mapping_key = f"{normalize_term(canonical)}|{normalize_term(scope)}"
                normalized_keys: List[str] = []
                for value in [canonical, *aliases_zh, official_en or "", *aliases_en]:
                    key = normalize_term(value)
                    if key and key not in normalized_keys:
                        normalized_keys.append(key)

                if mapping_key in key_rows:
                    row_errors.append(f"与第 {key_rows[mapping_key]} 行重复")
                elif canonical:
                    key_rows[mapping_key] = row_number

                for alias_key in normalized_keys:
                    alias_scope_key = (normalize_term(scope), alias_key)
                    previous = alias_targets.get(alias_scope_key)
                    if previous and previous[0] != canonical:
                        conflicts.append({
                            "row": row_number,
                            "previous_row": previous[1],
                            "scope": scope,
                            "alias": alias_key,
                            "standards": [previous[0], canonical],
                        })
                    elif canonical:
                        alias_targets[alias_scope_key] = (canonical, row_number)

                if row_errors:
                    errors.append({"row": row_number, "messages": row_errors})
                parsed.append({
                    "row": row_number,
                    "canonical_zh": canonical,
                    "aliases_zh": aliases_zh,
                    "official_en": official_en,
                    "aliases_en": aliases_en,
                    "scope": scope or None,
                    "note": _text(raw_payload["参考出处或补充说明"]) or None,
                    "raw_payload": raw_payload,
                    "normalized_keys": normalized_keys,
                    "mapping_key": mapping_key,
                    "errors": row_errors,
                })
        finally:
            workbook.close()

        return {
            "filename": filename,
            "sheet": sheet.title,
            "header_row": header_row,
            "total_rows": len(parsed),
            "skipped_rows": skipped,
            "valid_rows": len(parsed) - len(errors),
            "errors": errors,
            "conflicts": conflicts,
            "rows": parsed,
        }

    async def preview(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """只解析和校验，不写入任何数据库。"""
        result = self.parse(file_bytes, filename)
        result.pop("rows", None)
        result["can_import"] = not result["errors"] and not result["conflicts"]
        result["message"] = "校验通过，可以确认导入" if result["can_import"] else "请先修正错误或冲突"
        return result

    async def import_file(
        self,
        file_bytes: bytes,
        filename: str,
        *,
        upload_batch_id: Optional[str] = None,
        operator: Optional[str] = None,
    ) -> Dict[str, Any]:
        """通过全量预校验后，在一个事务中幂等写入用户表。"""
        result = self.parse(file_bytes, filename)
        if result["errors"] or result["conflicts"]:
            result.pop("rows", None)
            result["success_count"] = 0
            result["message"] = "校验未通过，未写入任何数据"
            return result

        batch_id = upload_batch_id or f"upload-{uuid.uuid4().hex}"
        created_count = 0
        updated_count = 0
        async with in_transaction() as connection:
            for row in result["rows"]:
                existing = await RagTermMappingSource.filter(
                    mapping_key=row["mapping_key"],
                ).using_db(connection).first()
                payload = {
                    "mapping_key": row["mapping_key"],
                    "canonical_zh": row["canonical_zh"],
                    "aliases_zh": row["aliases_zh"],
                    "official_en": row["official_en"],
                    "aliases_en": row["aliases_en"],
                    "scope": row["scope"],
                    "note": row["note"],
                    "raw_payload": row["raw_payload"],
                    "normalized_keys": row["normalized_keys"],
                    "status": "enabled",
                    "upload_batch_id": batch_id,
                    "source_file": filename,
                    "source_row": row["row"],
                    "version": (existing.version + 1) if existing else 1,
                    "updated_by": operator,
                }
                if existing:
                    for key, value in payload.items():
                        setattr(existing, key, value)
                    await existing.save(using_db=connection)
                    updated_count += 1
                else:
                    payload.update({
                        "mapping_id": f"src-{uuid.uuid4().hex}",
                        "created_by": operator,
                    })
                    await RagTermMappingSource.create(using_db=connection, **payload)
                    created_count += 1

        result.pop("rows", None)
        result.update({
            "upload_batch_id": batch_id,
            "success_count": created_count + updated_count,
            "created_count": created_count,
            "updated_count": updated_count,
            "message": "导入成功",
        })
        return result


term_mapping_import_service = TermMappingImportService()
