"""显式 --live：真实TCP/JWT/云MySQL/ZIP及前端API；仅替代沙箱边界。

不挂到生产服务，不创建测试控制接口，不写同事的任务表；结束清理本次唯一前缀数据。
"""
import argparse
import asyncio
from datetime import datetime, timedelta, timezone
import hashlib
import json
import logging
from pathlib import Path
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import traceback
import uuid

if sys.platform == 'win32':
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

import httpx
import jwt
from fastapi import FastAPI, Form, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from tortoise import Tortoise, connections
from tortoise.transactions import in_transaction
import uvicorn

from configure import PROJECT_CONFIG
from applications.wrq.auth import DependAuth, verify_token
from applications.wrq.database import wrq_connections
from applications.wrq.migrations.review_20261024 import connect, digest
from applications.wrq.migrations.review_20261026 import snapshot
from applications.wrq.migrations.schema_20261026 import TABLES
from applications.wrq.tests.debug_lifecycle_cases import cases, save_report

ROOT = Path(__file__).resolve().parents[3]
WORKSPACE = ROOT.parent
REPORT = WORKSPACE/'deliverables/开发者中心_调试后续流程测试结果_20261026.html'


class SandboxBoundary:
    """仅在本轮测试进程使用；其余既有任务仍只读同事表。"""
    def __init__(self, document):
        self.rows = {}
        self.document = document
        self.http_calls = []
        self.automatic_runs = set()

    def add(self, skill_id, status='queued', **changes):
        run_id = uuid.uuid4().hex
        self.rows[run_id] = {'run_id':run_id,'skill_id':skill_id,'status':status,
            'started_at':None,'finished_at':None,'expires_at':None,'error_message':None,
            'result_file_path':str(self.document),'result_file_name':self.document.name,**changes}
        return run_id

    def set(self, run_id, status, **changes):
        self.rows[run_id].update(status=status, **changes)
        if status == 'running':
            self.rows[run_id]['started_at'] = datetime.now(timezone.utc)
        if status in ('completed','failed'):
            self.rows[run_id]['finished_at'] = datetime.now(timezone.utc)

    async def get(self, run_id, skill_id):
        from applications.wrq.sandbox import SandboxResults
        row = self.rows.get(run_id)
        if row is None:
            return await SandboxResults().get(run_id, skill_id)
        if row['skill_id'] != skill_id:
            return None
        value = dict(row)
        value['status'] = {'completed':'PASSED','queued':'QUEUED','running':'RUNNING','failed':'FAILED'}.get(row['status'],row['status'].upper())
        return value

    def public(self, run_id):
        if run_id not in self.rows:
            raise HTTPException(404,'任务不存在')
        return {key:self.rows[run_id][key] for key in
                ('run_id','skill_id','status','started_at','finished_at','error_message')}


async def run():
    logging.getLogger('httpx').setLevel(logging.WARNING)
    report = {'phase':'执行中','cases':cases(),'runtime':{},'notes':[
        '沙箱任务状态和Excel为测试边界；开发者中心JWT、HTTP、云MySQL、构建及后续业务真实执行。',
        '使用临时WRQ测试服务，未绕过或改写运行中完整项目的公共鉴权；整项目鉴权单列X01。',
        '前端实际API函数通过TCP访问本轮服务；Node中仅浏览器文件缓存使用测试存储，未宣称完成浏览器UI/刷新验收。',
        '不修改同事代码或结果表，不创建用户表；只清理本次随机前缀的测试记录。',
        '首轮B16的用例前提有误：新版构建后已处于调试中；修正为比较迟到登记前后的当前版本/状态/通过标记均不变，而非断言构建前的草稿状态。']}
    index = {row['id']:row for row in report['cases']}
    def mark(case_id, ok, detail=''):
        index[case_id].update(status='PASSED' if ok else 'FAILED',detail=detail)
        save_report(REPORT,report)
        if not ok:
            raise AssertionError(case_id+': '+detail)
    def block(case_id, detail):
        index[case_id].update(status='BLOCKED',detail=detail)
    save_report(REPORT,report)
    temp_parent = (WORKSPACE/'.tmp').resolve()
    temp_parent.mkdir(exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix='wrq-debug-lifecycle-',dir=temp_parent)).resolve()
    prefix = 'WRQ调试验收-'+uuid.uuid4().hex[:12]+'-'
    previous_dir = PROJECT_CONFIG.SKILL_BUILD_DIR
    PROJECT_CONFIG.SKILL_BUILD_DIR = str(temporary/'builds')
    book = Workbook(); book.active.append(['用例编号','测试步骤','预期结果'])
    book.active.append(['MOCK-001','模型执行替代样例','用于下载后续链路校验'])
    document = temporary/'沙箱测试输出.xlsx'; book.save(document); book.close()
    boundary = SandboxBoundary(document)
    server = task = conn = sock = None
    original_crud = None
    fatal = None
    token_file = Path.home()/'.codex/private/wrq-interface-test-token.txt'
    token = token_file.read_text('utf8').strip()
    identity = verify_token(token)
    headers = {'token':token}
    try:
        conn = await connect()
        async with conn.cursor() as cursor:
            before = await snapshot(cursor)
            # 完整项目仍在运行，允许追加正常审计日志；已有审计记录必须逐值不变。
            await cursor.execute('SELECT * FROM keenrobot_audit')
            audit_before=list(await cursor.fetchall())
            audit_before.sort(key=lambda row:json.dumps(row,sort_keys=True,default=str))
            audit_ids={row['id'] for row in audit_before}
            before['keenrobot_audit']['hash']=digest(audit_before)
            before['keenrobot_audit']['count']=len(audit_before)
            await cursor.execute("SHOW SESSION STATUS LIKE 'Ssl_cipher'")
            report['runtime']['tls_cipher'] = (await cursor.fetchone())['Value']
        if not report['runtime']['tls_cipher']:
            raise RuntimeError('拒绝无TLS测试')
        await Tortoise.init(config={'connections':wrq_connections(PROJECT_CONFIG.DATABASE_CONNECTIONS),
            'apps':{'models':{'models':PROJECT_CONFIG.APPLICATIONS_MODELS,'default_connection':'default'}},
            'use_tz':False,'timezone':'Asia/Shanghai'})
        from applications.wrq.agent.models.agent_model import Skill, SkillVersion, SkillForm, SkillDebugRun, SkillPublication
        from applications.wrq.agent.services.developer_skill_crud import DeveloperSkillCrud
        from applications.wrq.agent.views import skills_router, developer_skill_view
        original_crud = developer_skill_view.crud
        developer_skill_view.crud = DeveloperSkillCrud(boundary)
        app = FastAPI(title='WRQ test-only sandbox boundary')
        app.include_router(skills_router,prefix='/skills')

        @app.post('/realtime-test-case-gen/runs',dependencies=[DependAuth])
        async def create_run(skill_id:str=Form(...)):
            boundary.http_calls.append({'method':'POST','path':'/realtime-test-case-gen/runs'})
            run_id=boundary.add(skill_id)
            boundary.automatic_runs.add(run_id)
            return boundary.public(run_id)

        @app.get('/realtime-test-case-gen/skills/{skill_id}/latest-run',dependencies=[DependAuth])
        async def latest_run(skill_id:str):
            rows=[r for r in boundary.rows.values() if r['skill_id']==skill_id]
            if not rows:
                raise HTTPException(404,'尚无任务')
            return boundary.public(rows[-1]['run_id'])

        @app.get('/realtime-test-case-gen/runs/{run_id}',dependencies=[DependAuth])
        async def get_run(run_id:str):
            if run_id in boundary.automatic_runs:
                status=boundary.rows[run_id]['status']
                if status in ('queued','running'):
                    boundary.set(run_id,'running' if status=='queued' else 'completed')
            return boundary.public(run_id)

        @app.get('/realtime-test-case-gen/runs/{run_id}/document',dependencies=[DependAuth])
        async def get_document(run_id:str):
            row=boundary.public(run_id)
            if row['status'] != 'completed':
                raise HTTPException(409,'任务未完成，暂无结果文档')
            return FileResponse(document,filename=document.name)

        sock=socket.socket(); sock.bind(('127.0.0.1',0)); sock.listen(128); sock.setblocking(False)
        base=f'http://127.0.0.1:{sock.getsockname()[1]}'
        server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=0,lifespan='off',access_log=False,log_level='warning'))
        task=asyncio.create_task(server.serve(sockets=[sock]))
        for _ in range(100):
            if server.started:
                break
            if task.done():
                await task
                raise RuntimeError('测试HTTP服务启动失败')
            await asyncio.sleep(.05)
        if not server.started:
            raise RuntimeError('测试HTTP服务启动超时')
        report['runtime'].update(test_service=base,database=PROJECT_CONFIG.DATABASE_NAME,
            token_source='既有私有Token文件，真实签名验证；不输出凭证',sandbox_mock=True,
            production_backend_changed=False,temporary_prefix=prefix)
        async with httpx.AsyncClient(base_url=base,headers=headers,trust_env=False,timeout=30) as client:
            async def call(method,path,expected=200,**kwargs):
                response=await client.request(method,path,**kwargs)
                if response.status_code!=expected:
                    raise AssertionError(f'{method} {path}: HTTP{response.status_code}, expected {expected}; {response.text[:250]}')
                payload=response.json()
                return payload.get('data')
            meta=await call('GET','/skills/developer/metadata')
            class_id=meta['skill_classes'][0]['id']
            async def create(label):
                return await call('POST','/skills/developer/',json={'name':prefix+label,'class_id':class_id,
                    'application_ids':['*'],'description':'沙箱后续真实链路验收','instructions':'生成正常、异常、边界用例',
                    'resource_ids':[],'uploaded_resources':[],'resource_usage':'','compatibility':''})
            def edit(row,**changes):
                return {key:row[key] for key in ('name','class_id','application_ids','description','instructions',
                    'resource_ids','uploaded_resources','resource_usage','compatibility','technical_name')}|{
                        'expected_revision':row['current_revision']}|changes
            async def build(row):
                source={'skill_id':row['id'],'skill_name':row['name'],'skill_class_id':row['class_id'],
                    'application':'*','description':row['description'],'instructions':row['instructions'],
                    'expected_version':row['version']}
                return await call('POST','/skills/form/builds',data={'source':json.dumps(source)})
            async def register(row,run_id,built,**changes):
                return await call('POST',f'/skills/developer/{row["id"]}/debug-runs',
                    json={'run_id':run_id,'build_id':built['build_id'],**changes})
            async def recover(sid,run_id,connection=client):
                response=await connection.get(f'/skills/developer/{sid}/debug-runs/{run_id}')
                if response.status_code!=200:
                    raise AssertionError(response.text[:200])
                return response.json()['data']
            async def skill_state(sid):
                return await Skill.get(id=sid)
            async def publications(sid):
                return await SkillPublication.filter(skill_id=sid).count()

            row=await create('主流程'); sid=row['id']; built=await build(row)
            version=await SkillVersion.get(id=row['skill_version_id'])
            mark('B01',Path(version.package_storage_key).is_file() and version.build_id==built['build_id']
                and await SkillForm.filter(skill_version_id=version.id).count()==1,
                f'技能{sid}；版本{version.id}；真实ZIP已落盘')
            run_id=boundary.add(sid)
            await register(row,run_id,built,conversation_input='测试登录',request_document_count=2)
            link=await SkillDebugRun.get(run_id=run_id)
            mark('B02',link.skill_version_id==version.id and link.build_id==built['build_id']
                 and link.resource_manifest_hash==version.manifest_hash,'HTTP登记成功；云库版本、构建和摘要一致')
            await register(row,run_id,built,conversation_input='重复提交不覆盖首次输入')
            mark('B03',await SkillDebugRun.filter(run_id=run_id).count()==1 and
                 (await SkillDebugRun.get(run_id=run_id)).conversation_input=='测试登录')
            async with httpx.AsyncClient(base_url=base,headers=headers,trust_env=False) as reconnected:
                restored=await recover(sid,run_id,reconnected)
            mark('B04',restored['skill_version_id']==version.id and restored['task']['status']=='QUEUED')
            boundary.set(run_id,'running')
            restored=await recover(sid,run_id)
            listed=next(r for r in await call('GET','/skills/developer/') if r['id']==sid)
            mark('B05',restored['task']['status']=='RUNNING' and listed['debug']['runId']==run_id
                 and listed['debug']['status']=='running' and restored['skill_version_id']==version.id)
            await call('POST',f'/skills/developer/{sid}/publish',409)
            mark('B06',not (await skill_state(sid)).is_enabled and await publications(sid)==0)
            boundary.set(run_id,'failed',error_message='测试沙箱429：模型额度不足')
            restored=await recover(sid,run_id)
            listed=next(r for r in await call('GET','/skills/developer/') if r['id']==sid)
            mark('B07',restored['task']['error_message']=='测试沙箱429：模型额度不足'
                 and listed['debug']['status']=='failed' and '429' in listed['debug']['output'])
            await register(row,run_id,built,status='passed')
            await call('POST',f'/skills/developer/{sid}/publish',409)
            mark('B08',not (await skill_state(sid)).is_enabled and await publications(sid)==0)
            retry=boundary.add(sid)
            await register(row,retry,built)
            mark('B09',await SkillDebugRun.filter(skill_id=sid).count()==2
                 and (await recover(sid,run_id))['task']['status']=='FAILED')
            boundary.set(retry,'running'); boundary.set(retry,'completed')
            row=await call('POST',f'/skills/developer/{sid}/publish')
            skill=await skill_state(sid)
            history=await SkillPublication.get(skill_id=sid)
            mark('B10',row['status']=='已上线' and skill.is_enabled and skill.published_at is not None
                and skill.last_debug_run_id==retry and history.skill_version_id==version.id)
            await call('POST',f'/skills/developer/{sid}/publish')
            mark('B11',await publications(sid)==1)
            off=await call('POST',f'/skills/developer/{sid}/offline')
            skill=await skill_state(sid)
            off_ok=off['status']=='已下线' and not skill.is_enabled and skill.offline_at is not None
            row=await call('POST',f'/skills/developer/{sid}/publish')
            mark('B12',off_ok and row['status']=='已上线' and await publications(sid)==3)
            original_instruction=(await SkillForm.get(skill_version_id=version.id)).instructions
            row=await call('PUT',f'/skills/developer/{sid}',json=edit(row,instructions='保存修改后的新版指令'))
            skill=await skill_state(sid)
            mark('B13',row['status']=='草稿' and row['skill_version_id']!=version.id and not skill.is_enabled
                and skill.last_debug_run_id is None and skill.last_debug_passed_at is None
                and (await SkillForm.get(skill_version_id=version.id)).instructions==original_instruction
                and await publications(sid)==3)
            restored=await recover(sid,retry)
            await call('POST',f'/skills/developer/{sid}/publish',409)
            mark('B14',not restored['is_current_version'] and restored['skill_version_id']==version.id)
            next_build=await build(row)
            await call('POST',f'/skills/developer/{sid}/debug-runs',409,json={'run_id':retry,'build_id':next_build['build_id']})
            mark('B15',(await SkillDebugRun.get(run_id=retry)).skill_version_id==version.id)
            late=boundary.add(sid)
            current_before=await skill_state(sid)
            await register(row,late,built)
            current_after=await skill_state(sid)
            mark('B16',(await SkillDebugRun.get(run_id=late)).skill_version_id==version.id
                and (current_before.current_version_id,current_before.display_status,current_before.last_debug_passed_at)
                    ==(current_after.current_version_id,current_after.display_status,current_after.last_debug_passed_at),
                f'旧任务关联版本{version.id}；当前版本{current_after.current_version_id}及状态{current_after.display_status}未改变')
            newer_run=boundary.add(sid)
            await register(row,newer_run,next_build)
            await call('POST',f'/skills/developer/{sid}/publish',409)
            mark('B17',not (await skill_state(sid)).is_enabled)
            boundary.set(newer_run,'running'); boundary.set(newer_run,'completed')
            row=await call('POST',f'/skills/developer/{sid}/publish')
            mark('B18',row['status']=='已上线' and (await skill_state(sid)).last_debug_run_id==newer_run
                and (await SkillDebugRun.get(run_id=newer_run)).skill_version_id==row['skill_version_id'])

            expired_row=await create('过期及异常'); expired_build=await build(expired_row)
            expired_sid=expired_row['id']
            expired_run=boundary.add(expired_sid,'completed',expires_at=datetime.now(timezone.utc)-timedelta(seconds=5))
            await register(expired_row,expired_run,expired_build)
            await call('POST',f'/skills/developer/{expired_sid}/publish',409)
            mark('B19',await publications(expired_sid)==0 and not (await skill_state(expired_sid)).is_enabled)
            before_links=await SkillDebugRun.all().count()
            invalids=[{'run_id':uuid.uuid4().hex,'build_id':expired_build['build_id']},
                {'run_id':retry,'build_id':expired_build['build_id']},
                {'run_id':expired_run,'build_id':'build-does-not-exist'}, {'run_id':expired_run}]
            for value in invalids:
                await call('POST',f'/skills/developer/{expired_sid}/debug-runs',409,json=value)
            mark('B20',await SkillDebugRun.all().count()==before_links)
            for value in ({'request_document_count':11},{'request_document_count':-1},{'run_id':''}):
                await call('POST',f'/skills/developer/{expired_sid}/debug-runs',422,
                    json={'run_id':expired_run,'build_id':expired_build['build_id'],**value})
            mark('B21',await SkillDebugRun.all().count()==before_links)
            conflict_run=boundary.add(sid)
            await call('POST',f'/skills/developer/{sid}/debug-runs',409,
                json={'run_id':conflict_run,'build_id':next_build['build_id'],'skill_version_id':version.id})
            mark('B22',not await SkillDebugRun.filter(run_id=conflict_run).exists())
            outsider=jwt.encode({'user_id':identity.id+100000,'username':'wrq-signed-permission-test',
                'state':0,'exp':datetime.now(timezone.utc)+timedelta(minutes=15)},
                PROJECT_CONFIG.AUTH_SECRET_KEY,algorithm=PROJECT_CONFIG.AUTH_JWT_ALGORITHM)
            original_skill=await Skill.filter(id=sid).values()
            for method,path,body in [
                ('GET',f'/skills/developer/{sid}/debug-runs/{retry}',None),
                ('POST',f'/skills/developer/{sid}/debug-runs',{'run_id':newer_run,'build_id':next_build['build_id']}),
                ('POST',f'/skills/developer/{sid}/publish',None),
                ('PUT',f'/skills/developer/{sid}',edit(row)),('DELETE',f'/skills/developer/{sid}',None)]:
                await call(method,path,404,headers={'token':outsider},**({'json':body} if body is not None else {}))
            mark('B23',original_skill==await Skill.filter(id=sid).values())
            concurrent_run=boundary.add(expired_sid)
            requests=await asyncio.gather(*[client.post(f'/skills/developer/{expired_sid}/debug-runs',
                json={'run_id':concurrent_run,'build_id':expired_build['build_id']}) for _ in range(2)])
            await register(expired_row,concurrent_run,expired_build)
            mark('B24',all(r.status_code in (200,409) for r in requests)
                and await SkillDebugRun.filter(run_id=concurrent_run).count()==1,
                f'并发HTTP状态{[r.status_code for r in requests]}，重试后仍只有一条')
            await call('DELETE',f'/skills/developer/{expired_sid}',400)
            saved=await call('PUT',f'/skills/developer/{expired_sid}',json=edit(expired_row,instructions='回草稿'))
            await call('DELETE',f'/skills/developer/{expired_sid}')
            reused=await call('POST','/skills/developer/',json=edit(saved)|{'expected_revision':None})
            mark('B25',reused['id']!=expired_sid and (await skill_state(expired_sid)).state==1
                and await SkillVersion.filter(skill_id=expired_sid).count()==2
                and await SkillDebugRun.filter(skill_id=expired_sid).count()==2)

            if 'keenrobot_user' not in before:
                block('X01','云库无keenrobot_user；完整项目公共鉴权依赖该表。未恢复表或替换公共鉴权。')
            else:
                from services import AuthControl
                await AuthControl.is_authed(token)
                mark('X01',True,'调用原项目公共鉴权验证既有Token成功')

            front=await create('前端API'); front_build=await build(front)
            queued=boundary.add(front['id']); completed=boundary.add(front['id'],'completed')
            config={'base':base,'skillId':front['id'],'buildId':front_build['build_id'],
                'versionId':front['skill_version_id'],'queuedRun':queued,'completedRun':completed,
                'excelHash':hashlib.sha256(document.read_bytes()).hexdigest(),
                'tokenFile':str(token_file),'reportFile':str(temporary/'frontend-results.json'),
                'bundleFile':str(temporary/'frontend-tests.mjs')}
            config_path=temporary/'frontend-config.json'
            config_path.write_text(json.dumps(config,ensure_ascii=False),encoding='utf8')
            # Windows云库TLS需要Selector loop；它不支持asyncio子进程，放到线程执行，
            # 不阻塞当前loop中的真实HTTP服务。
            process=await asyncio.to_thread(subprocess.run,
                ['node','scripts/test-debug-cloud.mjs',str(config_path)],
                cwd=str(WORKSPACE/'agent-template-frontend-master-latest'),
                stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=120)
            frontend_path=Path(config['reportFile'])
            if not frontend_path.exists():
                raise RuntimeError('前端测试未生成报告：'+process.stderr.decode('utf8',errors='replace')[-1200:])
            frontend=json.loads(frontend_path.read_text('utf8'))
            report['runtime']['frontend_http_calls']=frontend['requests']
            for result in frontend['cases']:
                index[result['id']].update(status=result['status'],detail=result['detail'])
            # 独立SQL核验前端结果，不能只相信Node内存断言。
            early=await SkillDebugRun.filter(run_id=queued).exists()
            if not early:
                index['F02'].update(status='FAILED',detail='前端running状态直接返回；云库中没有该queued任务关联')
            front_skill=await skill_state(front['id'])
            sql_ok=(await SkillDebugRun.filter(run_id=completed,skill_version_id=front['skill_version_id']).exists()
                    and front_skill.current_revision==2 and front_skill.display_status=='DRAFT'
                    and not front_skill.is_enabled and await publications(front['id'])==3)
            if not sql_ok:
                index['F07'].update(status='FAILED',detail='前端后续流程的独立SQL复核未通过')
            report['runtime']['frontend_sql_verified']=sql_ok
            report['runtime']['frontend_exit_code']=process.returncode

    except Exception as exc:
        location=traceback.extract_tb(exc.__traceback__)[-1]
        fatal={'type':type(exc).__name__,'message':str(exc),'where':Path(location.filename).name+':'+str(location.lineno)}
    finally:
        if server:
            server.should_exit=True
        if task:
            try:
                await asyncio.wait_for(task,timeout=10)
            except (Exception,asyncio.CancelledError):
                task.cancel()
        if sock:
            sock.close()
        if original_crud is not None:
            developer_skill_view.crud=original_crud
        try:
            if Tortoise.apps:
                owned=await Skill.filter(name__startswith=prefix,user_id=str(identity.id)).values_list('id',flat=True)
                async with in_transaction() as db:
                    if owned:
                        placeholders=','.join(['%s']*len(owned))
                        await db.execute_query(f'UPDATE keenrobot_skills SET current_version_id=NULL WHERE id IN ({placeholders})',owned)
                        for table in ('keenrobot_skill_resource_bindings','keenrobot_skill_debug_runs','keenrobot_skill_publications'):
                            await db.execute_query(f'DELETE FROM {table} WHERE skill_id IN ({placeholders})',owned)
                        await db.execute_query(f'DELETE FROM keenrobot_skill_forms WHERE skill_version_id IN '
                            f'(SELECT id FROM keenrobot_skill_versions WHERE skill_id IN ({placeholders}))',owned)
                        await db.execute_query(f'DELETE FROM keenrobot_skill_versions WHERE skill_id IN ({placeholders})',owned)
                        await db.execute_query(f'DELETE FROM keenrobot_skills WHERE id IN ({placeholders})',owned)
                report['runtime']['cleaned_skill_count']=len(owned)
            if conn:
                async with conn.cursor() as cursor:
                    after=await snapshot(cursor)
                    await cursor.execute('SELECT * FROM keenrobot_audit')
                    current_audit=list(await cursor.fetchall())
                    original_audit=[row for row in current_audit if row['id'] in audit_ids]
                    original_audit.sort(key=lambda row:json.dumps(row,sort_keys=True,default=str))
                # 自增计数自然前进；其余结构、所有原有行须保持。
                normalize=lambda ddl:re.sub(r' AUTO_INCREMENT=\d+','',ddl)
                protected={name:{'rows_unchanged':before[name]['hash']==after[name]['hash'],
                    'structure_unchanged':normalize(before[name]['ddl'])==normalize(after[name]['ddl'])}
                    for name in before if name in after}
                report['runtime']['preservation_checks']=protected
                audit_unchanged=digest(original_audit)==before['keenrobot_audit']['hash']
                report['runtime']['audit_preservation']={'original_rows_unchanged':audit_unchanged,
                    'appended_rows':len(current_audit)-len(audit_before),'table_structure_unchanged':protected['keenrobot_audit']['structure_unchanged']}
                remaining=all(all(value.values()) for name,value in protected.items() if name!='keenrobot_audit')
                mark('D01',before.keys()==after.keys() and remaining and audit_unchanged and protected['keenrobot_audit']['structure_unchanged'],
                    '本轮测试记录已清理；所有原业务行、原审计行及同事表保持不变。允许运行中的整项目追加正常审计日志。')
        except Exception as exc:
            fatal=fatal or {'type':type(exc).__name__,'message':'清理核验失败：'+str(exc)}
            index['D01'].update(status='FAILED',detail=str(exc))
        await Tortoise.close_connections()
        if conn:
            conn.close()
        PROJECT_CONFIG.SKILL_BUILD_DIR=previous_dir
        if temporary.parent!=temp_parent or not temporary.name.startswith('wrq-debug-lifecycle-'):
            raise RuntimeError('临时目录越界，拒绝删除')
        shutil.rmtree(temporary)
        for case in report['cases']:
            if case['status']=='PENDING':
                case.update(status='BLOCKED',detail='前置流程未完成：'+str(fatal))
        report.update(phase='测试结束',fatal=fatal)
        report['runtime']['cleanup_complete']=index['D01']['status']=='PASSED'
        save_report(REPORT,report)
    counts={state:sum(case['status']==state for case in report['cases']) for state in ('PASSED','FAILED','BLOCKED')}
    print(json.dumps({'counts':counts,'fatal':fatal,'report':str(REPORT)},ensure_ascii=False))
    if fatal or counts['FAILED'] or counts['BLOCKED']:
        raise SystemExit(1)


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--live',action='store_true')
    if not parser.parse_args().live:
        parser.error('须显式--live才会创建本轮云库测试记录；结束后清理')
    asyncio.run(run())
