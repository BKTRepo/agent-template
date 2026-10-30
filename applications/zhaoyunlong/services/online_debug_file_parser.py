"""在线调试上传文件的受限文本提取。"""

from __future__ import annotations

import io
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Final


MAX_UPLOAD_BYTES: Final = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS: Final = 500_000
TEXT_SUFFIXES: Final = {".md", ".txt", ".json"}
FILE_SUFFIXES: Final = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx"}


class FileParseError(ValueError):
    """上传文件无法安全提取为模型可读文本。"""


def extract_uploaded_file(filename: str, content: bytes) -> dict[str, str]:
    """校验并提取前端上传的六种允许格式文件。"""
    safe_name = _safe_filename(filename)
    suffix = Path(safe_name).suffix.lower()
    if suffix not in FILE_SUFFIXES:
        raise FileParseError("仅支持 PDF、DOCX、Markdown、TXT、XLSX、JSON 文件")
    return _extract_file_content(safe_name, suffix, content)


def extract_mounted_resource(filename: str, content: bytes) -> dict[str, str]:
    """提取已挂载资源文本，不重复执行前端上传文件的格式白名单校验。"""
    safe_name = _safe_filename(filename)
    suffix = Path(safe_name).suffix.lower()
    return _extract_file_content(safe_name, suffix, content, allow_other_text=True)


def _safe_filename(filename: str) -> str:
    """仅接受文件名，拒绝 Windows 与 POSIX 两种路径分隔符的路径穿越。"""
    safe_name = filename.strip()
    if (
        not safe_name
        or PurePosixPath(safe_name).name != safe_name
        or PureWindowsPath(safe_name).name != safe_name
    ):
        raise FileParseError("文件名不合法")
    return safe_name


def _extract_file_content(
    safe_name: str,
    suffix: str,
    content: bytes,
    *,
    allow_other_text: bool = False,
) -> dict[str, str]:
    """将受支持文件或已解析资源转换为容器可读取的文本。"""
    if not content:
        raise FileParseError(f"上传文件为空：{safe_name}")
    if len(content) > MAX_UPLOAD_BYTES:
        raise FileParseError(f"上传文件超过 20 MB：{safe_name}")

    if suffix in TEXT_SUFFIXES:
        text = _decode_text(content, safe_name)
        sandbox_name = safe_name
    elif suffix == ".pdf":
        text = _extract_pdf_text(content, safe_name)
        sandbox_name = f"{safe_name}.md"
    elif suffix == ".docx":
        text = _extract_docx_text(content, safe_name)
        sandbox_name = f"{safe_name}.md"
    elif suffix == ".xlsx":
        text = _extract_xlsx_text(content, safe_name)
        sandbox_name = f"{safe_name}.md"
    elif allow_other_text:
        # 已挂载资源的格式已由资源系统确认；此处只尝试读取其文本内容，不再使用格式白名单拒绝。
        text = _decode_text(content, safe_name)
        sandbox_name = safe_name
    else:
        raise FileParseError("仅支持 PDF、DOCX、Markdown、TXT、XLSX、JSON 文件")

    if not text.strip():
        raise FileParseError(f"未能从文件提取可读文本：{safe_name}")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise FileParseError(f"文件提取后的文本超过 500000 个字符：{safe_name}")
    return {"filename": sandbox_name, "content": f"# 上传文件：{safe_name}\n\n{text.strip()}\n"}


def _decode_text(content: bytes, filename: str) -> str:
    """读取 UTF-8 文本文件。"""
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileParseError(f"文本文件必须为 UTF-8 编码：{filename}") from exc


def _extract_pdf_text(content: bytes, filename: str) -> str:
    """读取 PDF 的文字层；扫描件需要由前端或独立 OCR 流程先行识别。"""
    try:
        import fitz
        document = fitz.open(stream=content, filetype="pdf")
        try:
            return "\n\n".join(page.get_text() for page in document)
        finally:
            document.close()
    except ImportError as exc:
        raise FileParseError("服务端缺少 PDF 解析依赖 PyMuPDF") from exc
    except Exception as exc:
        raise FileParseError(f"PDF 解析失败：{filename}") from exc


def _extract_docx_text(content: bytes, filename: str) -> str:
    """读取 DOCX 的段落和表格文字。"""
    try:
        from docx import Document
        document = Document(io.BytesIO(content))
        parts = [item.text for item in document.paragraphs if item.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ImportError as exc:
        raise FileParseError("服务端缺少 Word 解析依赖 python-docx") from exc
    except Exception as exc:
        raise FileParseError(f"DOCX 解析失败：{filename}") from exc


def _extract_xlsx_text(content: bytes, filename: str) -> str:
    """读取 XLSX 单元格值，不计算公式。"""
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        parts: list[str] = []
        cell_count = 0
        for sheet in workbook.worksheets:
            parts.append(f"## 工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cell_count += len(row)
                if cell_count > 100_000:
                    raise FileParseError(f"XLSX 单元格数量超过限制：{filename}")
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    parts.append(" | ".join(values))
        return "\n".join(parts)
    except FileParseError:
        raise
    except ImportError as exc:
        raise FileParseError("服务端缺少 Excel 解析依赖 openpyxl") from exc
    except Exception as exc:
        raise FileParseError(f"XLSX 解析失败：{filename}") from exc
