"""定稿知识库：真实 ORM/HTTP/文件解析；向量服务用显式测试替身隔离外部模型。"""
import hashlib
import io
import json
from pathlib import Path

from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
import pytest
import pytest_asyncio

from applications.agent.views import skills_router
from applications.knowledge_base.views import knowledge_router
from applications.knowledge_base.models.knowledge_base_model import Document, DocumentChunk, KnowledgeBase
from applications.knowledge_base.schemas.platform_schema import KnowledgeForm
from applications.knowledge_base.services.platform_service import KnowledgeService, get_knowledge_service, KnowledgeError
from applications.knowledge_base.services.document_pipeline import DocumentPipeline
from applications.knowledge_base.services.document_loader import load_document_pages
from applications.user.models.user_model import User
from configure import PROJECT_CONFIG
from enums import DocumentStatus
from test_skill_package_api import token


FORM = {"knowledge_name": "转账业务知识", "description": "转账功能测试的业务规则", "tag": "custom",
        "application": "核心系统", "is_public": True}


class TestIndex:
    __test__ = False
    def __init__(self):
        self.rows, self.calls, self.fail = {}, [], False
    async def embed(self, texts):
        self.calls.append(texts)
        if self.fail:
            raise RuntimeError("isolated embedding failure")
        return [[0.2, 0.5, 0.3] for _ in texts]
    async def put(self, kb_id, rows):
        for row in rows:
            self.rows[row["vector_id"]] = {**row, "kb_id": kb_id}
        return [row["vector_id"] for row in rows]
    async def remove(self, ids):
        for key in ids:
            self.rows.pop(key, None)


@pytest_asyncio.fixture
async def kb_setup(database, tmp_path, monkeypatch):
    monkeypatch.setattr(PROJECT_CONFIG, "OUTPUT_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(PROJECT_CONFIG, "SKILL_BUILD_DIR", str(tmp_path / "builds"))
    index = TestIndex()
    service = KnowledgeService(DocumentPipeline(index))
    owner = await User.create(username="owner", alias="王", email="owner@example.com", password="unused")
    other = await User.create(username="other", alias="赵", email="other@example.com", password="unused", is_superuser=True)
    app = FastAPI()
    app.include_router(knowledge_router, prefix="/knowledge-bases")
    app.include_router(skills_router, prefix="/skills")
    app.dependency_overrides[get_knowledge_service] = lambda: service
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, owner, other, service, index


async def save(client, user, data=None, files=None, kb_id=None):
    method = client.put if kb_id else client.post
    return await method("/knowledge-bases/" + (kb_id or ""), headers=token(user),
        data={"source": json.dumps(data or FORM)}, files=files if files is not None else [("files", ("rules.md", "# 转账规则\n\n余额必须充足。".encode()))])


@pytest.mark.asyncio
async def test_create_real_files_search_preview_download_and_history(kb_setup):
    client, owner, _, _, index = kb_setup
    response = await save(client, owner)
    assert response.status_code == 200, response.text
    kb = response.json()["data"]
    assert kb["version"] == "V1.0.0" and kb["user_id"] == owner.id
    detail = (await client.get(f"/knowledge-bases/{kb['id']}", headers=token(owner))).json()["data"]
    assert detail["status"] == "completed" and detail["document_count"] == 1
    doc = detail["documents"][0]
    assert doc["status"] == "completed" and doc["chunk_count"] > 0 and index.rows
    preview = (await client.get(f"/knowledge-bases/{kb['id']}/documents/{doc['id']}/preview", headers=token(owner))).json()["data"]
    assert "余额必须充足" in preview["markdown"]
    download = await client.get(f"/knowledge-bases/{kb['id']}/documents/{doc['id']}/download", headers=token(owner))
    assert download.content == "# 转账规则\n\n余额必须充足。".encode()
    rows = (await client.get("/knowledge-bases/?search=转账", headers=token(owner))).json()
    assert rows["total"] == 1
    assert (await client.get("/knowledge-bases/?search=不存在", headers=token(owner))).json()["total"] == 0
    chunks = (await client.get(f"/knowledge-bases/{kb['id']}/chunks", headers=token(owner))).json()["data"]
    assert chunks[0]["source_metadata"]["section_title"] == "转账规则"


@pytest.mark.asyncio
async def test_personal_permissions_even_superuser_cannot_read_private_or_edit_others(kb_setup):
    client, owner, other, _, _ = kb_setup
    private = (await save(client, owner, {**FORM, "is_public": False})).json()["data"]
    assert (await client.get("/knowledge-bases/", headers=token(other))).json()["total"] == 0
    assert (await client.get(f"/knowledge-bases/{private['id']}", headers=token(other))).status_code == 404
    public = (await save(client, owner, {**FORM, "knowledge_name": "公开业务"})).json()["data"]
    assert (await client.get(f"/knowledge-bases/{public['id']}", headers=token(other))).json()["data"]["can_edit"] is False
    assert (await save(client, other, {**FORM, "change_summary": "越权修改"}, files=[], kb_id=public["id"])).status_code == 403
    assert (await client.post(f"/knowledge-bases/{public['id']}/retry", headers=token(other))).status_code == 403


@pytest.mark.asyncio
async def test_required_files_name_uniqueness_and_bad_files_cleaned(kb_setup):
    client, owner, _, _, _ = kb_setup
    assert (await save(client, owner, files=[])).status_code == 400
    assert await KnowledgeBase.all().count() == 0
    first = await save(client, owner, {**FORM, "knowledge_name": "ＡＢＣ"})
    assert first.json()["status"] == "success"
    duplicate = await save(client, owner, {**FORM, "knowledge_name": "abc"})
    assert duplicate.status_code == 400
    for filename, content in [("../a.md", b"x"), ("a.exe", b"x"), ("empty.txt", b"")]:
        assert (await save(client, owner, {**FORM, "knowledge_name": "新库"}, [("files", (filename, content))])).status_code == 400
    assert await KnowledgeBase.all().count() == 1
    assert len(list(PROJECT_CONFIG.upload_path.glob("*/*"))) == 1
    forged = await save(client, owner, {**FORM, "user_id": 999})
    assert forged.status_code == 400


@pytest.mark.asyncio
async def test_metadata_edit_does_not_reprocess_and_document_edit_only_changes_one(kb_setup):
    client, owner, _, _, index = kb_setup
    kb = (await save(client, owner, files=[("files", ("one.md", b"# One\nfirst rule")), ("files", ("two.txt", b"second rule"))])).json()["data"]
    detail = (await client.get(f"/knowledge-bases/{kb['id']}", headers=token(owner))).json()["data"]
    calls = len(index.calls)
    result = await save(client, owner, {**FORM, "description": "修改描述", "change_summary": "修订说明", "expected_version": "V1.0.0"}, files=[], kb_id=kb["id"])
    assert result.json()["data"]["version"] == "V1.0.1"
    assert len(index.calls) == calls
    doc = next(d for d in detail["documents"] if d["filename"] == "one.md")
    edit = await client.put(f"/knowledge-bases/{kb['id']}/documents/{doc['id']}/content", headers=token(owner),
        json={"content": "# One\nnew rule", "change_summary": "更新一号文档", "expected_generation": doc["generation"]})
    assert edit.json()["data"]["version"] == "V1.0.2", edit.text
    assert len(index.calls) == calls + 1
    assert len(index.rows) == await DocumentChunk.all().count()
    again = await client.put(f"/knowledge-bases/{kb['id']}/documents/{doc['id']}/content", headers=token(owner),
        json={"content": "stale", "change_summary": "过期编辑", "expected_generation": 1})
    assert again.status_code == 409
    current = (await client.get(f"/knowledge-bases/{kb['id']}", headers=token(owner))).json()["data"]
    assert [h["version"] for h in current["history"]] == ["V1.0.0", "V1.0.1", "V1.0.2"]


@pytest.mark.asyncio
async def test_embedding_failure_keeps_original_and_previous_index_then_retry(kb_setup):
    client, owner, _, _, index = kb_setup
    kb = (await save(client, owner)).json()["data"]
    old_ids = set(index.rows)
    index.fail = True
    updated = await save(client, owner, {**FORM, "change_summary": "更新规则"}, [("files", ("rules.md", b"# Updated\nnew rules"))], kb["id"])
    assert updated.status_code == 200
    detail = (await client.get(f"/knowledge-bases/{kb['id']}", headers=token(owner))).json()["data"]
    assert detail["status"] == "failed" and set(index.rows) == old_ids
    doc = await Document.get(id=detail["documents"][0]["id"])
    assert Path(doc.file_path).read_bytes() == b"# Updated\nnew rules" and "Updated" in doc.markdown
    assert set(await DocumentChunk.all().values_list("chroma_id", flat=True)) == old_ids
    index.fail = False
    assert (await client.post(f"/knowledge-bases/{kb['id']}/retry", headers=token(owner))).status_code == 200
    detail = (await client.get(f"/knowledge-bases/{kb['id']}", headers=token(owner))).json()["data"]
    assert detail["status"] == "completed" and not old_ids.intersection(index.rows)


@pytest.mark.asyncio
async def test_knowledge_has_no_skill_size_limit_and_unchanged_reupload_is_not_reindexed(kb_setup, monkeypatch):
    client, owner, _, _, index = kb_setup
    monkeypatch.setattr(PROJECT_CONFIG, "SKILL_RESOURCE_MAX_BYTES", 2)
    kb = (await save(client, owner)).json()["data"]
    assert kb["document_count"] == 1
    calls = len(index.calls)
    result = await save(client, owner, {**FORM, "change_summary": "同内容重传"}, kb_id=kb["id"])
    assert result.status_code == 200 and len(index.calls) == calls
    assert len(list(PROJECT_CONFIG.upload_path.glob("*/*"))) == 1


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pdf", "md", "txt"])
def test_real_document_parsers(tmp_path, kind):
    path = tmp_path / ("sample." + kind)
    if kind == "docx":
        from docx import Document as Word
        doc = Word(); doc.add_heading("Rules", 1); doc.add_paragraph("balance must be sufficient")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "minimum amount"; doc.save(path)
    elif kind == "xlsx":
        from openpyxl import Workbook
        book = Workbook(); book.active.append(["balance must be sufficient"]); book.save(path)
    elif kind == "pdf":
        import fitz
        doc = fitz.open(); doc.new_page().insert_text((50, 50), "balance must be sufficient"); doc.save(path); doc.close()
    else:
        path.write_text("balance must be sufficient", encoding="utf-8")
    text = "\n".join(p.page_content for p in load_document_pages(str(path), kind))
    assert "balance must be sufficient" in text


@pytest.mark.asyncio
async def test_remove_file_reclaims_index_and_rejects_empty_knowledge(kb_setup):
    client, owner, _, _, index = kb_setup
    kb = (await save(client, owner, files=[("files", ("one.md", b"first")), ("files", ("two.txt", b"second"))])).json()["data"]
    doc_id = kb["documents"][0]["id"]
    result = await save(client, owner, {**FORM, "change_summary": "删除第二份文件", "retained_document_ids": [doc_id]}, files=[], kb_id=kb["id"])
    assert result.status_code == 200
    assert await Document.all().count() == 1 and len(index.rows) == 1
    empty = await save(client, owner, {**FORM, "change_summary": "删除全部", "retained_document_ids": []}, files=[], kb_id=kb["id"])
    assert empty.status_code == 400 and await Document.all().count() == 1


@pytest.mark.asyncio
async def test_crashed_job_recovery_and_failed_index_cleanup_are_durable(kb_setup):
    from datetime import datetime, timedelta, timezone
    client, owner, _, service, index = kb_setup
    kb = (await save(client, owner)).json()["data"]
    doc = await Document.get(id=kb["documents"][0]["id"])
    orphan = "orphan-from-interrupted-vector-write"
    index.rows[orphan] = {"doc_id": doc.id}
    doc.cleanup_vector_ids = [orphan]
    doc.status = DocumentStatus.PROCESSING
    doc.processing_started_time = datetime.now(timezone.utc) - timedelta(hours=1)
    await doc.save()
    jobs = await service.retry(kb["id"], owner)
    assert jobs == [doc.id]
    await service.process(doc.id)
    await doc.refresh_from_db()
    assert doc.status == DocumentStatus.COMPLETED and not doc.cleanup_vector_ids
    assert orphan not in index.rows and len(index.rows) == await DocumentChunk.all().count()
    calls = len(index.calls)
    await service.process(doc.id)
    assert len(index.calls) == calls  # 已完成任务的重复回调不重复嵌入。


@pytest.mark.asyncio
async def test_name_backfill_audits_collisions_before_writing(kb_setup):
    from applications.knowledge_base.services.name_migration import migrate_knowledge_names
    _, owner, _, _, _ = kb_setup
    one = await KnowledgeBase.create(user_id=owner.id, knowledge_name="ＡＢＣ")
    two = await KnowledgeBase.create(user_id=owner.id, knowledge_name="abc")
    report = await migrate_knowledge_names(apply=True)
    assert not report["ready"] and not report["applied"]
    await one.refresh_from_db()
    assert one.name_key is None
    two.state = 1
    await two.save()
    assert (await migrate_knowledge_names(apply=True))["applied"]
